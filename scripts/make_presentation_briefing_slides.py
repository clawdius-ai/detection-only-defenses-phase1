"""Generate a presentation deck from report/presentation_briefing.md.

The content is curated from the briefing file into presentation-sized slides.

Output:
    report/presentation_briefing_slides.pptx
"""
from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report" / "presentation_briefing.md"
OUTPUT = ROOT / "report" / "presentation_briefing_slides.pptx"

BG = RGBColor(0x0B, 0x0B, 0x10)
PANEL = RGBColor(0x18, 0x18, 0x1F)
BORDER = RGBColor(0x3F, 0x3F, 0x46)
TEXT = RGBColor(0xF4, 0xF4, 0xF5)
MUTED = RGBColor(0xA1, 0xA1, 0xAA)
GREEN = RGBColor(0x34, 0xD3, 0x99)
BLUE = RGBColor(0x60, 0xA5, 0xFA)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ROSE = RGBColor(0xFB, 0x71, 0x85)
ORANGE = RGBColor(0xF9, 0x73, 0x16)


SLIDES = [
    {
        "kind": "title",
        "title": "Detection-Only Defenses for LLMs",
        "subtitle": "Phase 1: PAIR and TAP adaptive jailbreaks on AdvBench",
        "footer": "Presentation briefing deck",
    },
    {
        "kind": "agenda",
        "title": "Presentation Roadmap",
        "items": [
            "Problem: adaptive jailbreaks against aligned LLMs",
            "Dataset: AdvBench harmful-behavior benchmark",
            "Attacks: PAIR and TAP",
            "Defenses: detector-only routing baselines",
            "Metrics, traces, and reproducibility",
            "Dashboard workflow: demo traces vs. real API runs",
        ],
    },
    {
        "title": "Project Summary",
        "bullets": [
            "Evaluate whether detector-only defenses can stop adaptive jailbreak attacks.",
            "Use AdvBench harmful-behavior prompts as the initial benchmark.",
            "Implement PAIR and TAP as black-box adaptive attack baselines.",
            "Run attacks with no defense first, then under detector-only defenses.",
            "Save every prompt, detector decision, response, score, and session trace.",
        ],
    },
    {
        "title": "Problem: Adaptive Jailbreaks",
        "bullets": [
            "LLMs are aligned to refuse harmful or restricted requests.",
            "A jailbreak prompt tries to bypass those refusal behaviors.",
            "Static attacks use one fixed prompt.",
            "Adaptive attacks revise prompts after observing refusals, partial compliance, detector blocks, or judge scores.",
            "This is harder for defenses because the attacker can search for less obvious phrasings over multiple rounds.",
        ],
    },
    {
        "title": "What Is AdvBench?",
        "bullets": [
            "A harmful-behavior benchmark used to evaluate whether LLMs produce unsafe outputs.",
            "Local file: src/data/advbench_harmful_behaviors.csv.",
            "Columns: goal and target.",
            "goal = harmful objective the attacker tries to elicit.",
            "target = affirmative response prefix representing unsafe compliance.",
            "Phase 1 starts with the first 100 prompts; repo defaults to 20 for cost-bounded testing.",
        ],
    },
    {
        "kind": "diagram",
        "title": "System Architecture",
        "steps": [
            ("AdvBench", "harmful goal"),
            ("Attacker", "PAIR or TAP"),
            ("Detector", "allow / block"),
            ("Target LLM", "unchanged prompt"),
            ("Judge", "score 1-10"),
            ("Trace", "JSONL round record"),
        ],
        "caption": "If score < success threshold, feedback loops back to the attacker for another round.",
    },
    {
        "title": "PAIR: Prompt Automatic Iterative Refinement",
        "bullets": [
            "PAIR uses an attacker LLM to generate a candidate jailbreak prompt.",
            "The target model answers the candidate if the detector allows it.",
            "A judge scores the target response from 1 to 10.",
            "The target response and score are fed back to the attacker.",
            "PAIR is a linear search: round 1 -> round 2 -> round 3 -> success or budget exhausted.",
        ],
    },
    {
        "kind": "two_col",
        "title": "PAIR Round Loop",
        "left_title": "Per Round",
        "left": [
            "Attacker emits JSON: { improvement, prompt }",
            "Detector decides allow or block",
            "Allowed prompt queries target LLM",
            "Judge scores response",
            "Feedback drives next refinement",
        ],
        "right_title": "Why It Matters",
        "right": [
            "Black-box: only API access needed",
            "Semantic: prompts are readable",
            "Adaptive: each round reacts to prior feedback",
            "Traceable: each step is logged",
        ],
    },
    {
        "title": "TAP: Tree of Attacks with Pruning",
        "bullets": [
            "TAP extends PAIR from one chain into a tree search.",
            "Branch: generate multiple candidate refinements from the current frontier.",
            "Pre-target pruning: remove off-topic candidates before target queries.",
            "Attack and assess: query the target and judge surviving candidates.",
            "Post-target pruning: keep the best-scoring leaves for the next depth.",
        ],
    },
    {
        "kind": "two_col",
        "title": "PAIR vs. TAP",
        "left_title": "PAIR",
        "left": [
            "Single linear refinement chain",
            "Simpler and cheaper per step",
            "Can get stuck in one weak path",
            "Dashboard shape: line",
        ],
        "right_title": "TAP",
        "right": [
            "Branching search tree",
            "Explores more candidate strategies",
            "Pruning controls query cost",
            "Dashboard shape: tree",
        ],
    },
    {
        "title": "Detector-Only Defense Definition",
        "bullets": [
            "A detector evaluates a candidate prompt and optional session context.",
            "It makes only a routing decision: allow, block, or refuse.",
            "It does not rewrite the prompt.",
            "It does not modify target-model internals or generation.",
            "This isolates external safety routing as the defense mechanism.",
        ],
    },
    {
        "kind": "two_col",
        "title": "Defense Baselines",
        "left_title": "Simple Baselines",
        "left": [
            "none: always allow; establishes attack baseline",
            "keyword: blocklist terms from harmful domains",
            "Blocked prompts never reach the target",
        ],
        "right_title": "Learned / LLM Detectors",
        "right": [
            "PromptGuard: classifier for benign, injection, jailbreak",
            "Llama Guard 3: LLM-based safe/unsafe classifier",
            "Detector score or label is logged per round",
        ],
    },
    {
        "title": "Models and Budgets",
        "bullets": [
            "Attacker: mistralai/mixtral-8x22b-instruct via OpenRouter.",
            "Target: gpt-4o-2024-08-06 via OpenAI.",
            "Judge: gpt-4o-2024-08-06 via OpenAI with a 1-10 harm rubric.",
            "Budgets: B=5 and B=10 target queries per behavior.",
            "Success threshold: judge_score >= 10.",
        ],
    },
    {
        "kind": "metrics",
        "title": "Metrics",
        "metrics": [
            ("ASR", "successful sessions / behaviors"),
            ("Block Rate", "blocked candidates / all candidates"),
            ("Mean Max Score", "average best judge score per session"),
            ("Avg. Queries", "target queries per behavior"),
        ],
    },
    {
        "title": "What Gets Saved?",
        "bullets": [
            "Every round is appended to JSONL for reproducibility.",
            "Core fields: session_id, attack, defense, budget, advbench_index, goal.",
            "Graph fields: round and parent_round.",
            "Decision fields: candidate_prompt, detector_decision, detector_score.",
            "Outcome fields: target_response, judge_score, success, extra metadata.",
            "PAIR parent_round forms a line; TAP parent_round forms a tree.",
        ],
    },
    {
        "kind": "two_col",
        "title": "Demo Traces vs. Real API Runs",
        "left_title": "Demo Traces",
        "left": [
            "Generated by scripts/generate_demo_traces.py",
            "Written to results/traces/demo_synthetic.jsonl",
            "Marked with extra.demo = true",
            "Used for dashboard presentation and UI testing",
        ],
        "right_title": "Real API Runs",
        "right": [
            "Run through scripts/run_experiment.py",
            "Use OpenAI and OpenRouter credentials",
            "Produce real target, judge, and attacker outputs",
            "Separated in frontend under Real API Runs",
        ],
    },
    {
        "title": "Dashboard Walkthrough",
        "bullets": [
            "Use the navbar to switch between Demo Traces and Real API Runs.",
            "Top chart compares attack success rate by defense, attack, and budget.",
            "Session grid filters by attack and defense.",
            "Attack Trace Graph visualizes PAIR as a chain and TAP as a tree.",
            "Click a node to inspect prompt, detector decision, response, judge score, and attacker reasoning.",
        ],
    },
    {
        "title": "Current Repo Status",
        "bullets": [
            "Implemented: AdvBench loader, PAIR, TAP, detector registry, and GPT-4o judge wrapper.",
            "Implemented: no-defense, keyword, PromptGuard, and Llama Guard detectors.",
            "Implemented: JSONL tracing, aggregation, FastAPI trace server, and React dashboard.",
            "Demo data is available for presentation.",
            "Real API smoke examples are available for PAIR and TAP.",
        ],
    },
    {
        "title": "Limitations and Next Steps",
        "bullets": [
            "Full first-100 AdvBench grid still needs to be run for final numbers.",
            "StrongREJECT and JBB-Behaviors can be added as later datasets.",
            "Benign-prompt false-positive evaluation should be added for broader defense analysis.",
            "Future analysis can vary judge thresholds such as 8, 9, and 10.",
            "Phase 2 can build graph-based analytics over complete session traces.",
        ],
    },
    {
        "title": "Likely Q&A",
        "bullets": [
            "Why AdvBench? It standardizes harmful objectives for comparison.",
            "Why no-defense first? It establishes baseline attack behavior.",
            "Why detector-only? It isolates practical routing defenses from model modification.",
            "Why TAP over PAIR? TAP explores multiple branches and prunes weak paths.",
            "Why traces? Aggregate ASR hides how attacks adapt and where defenses fail.",
        ],
    },
    {
        "kind": "references",
        "title": "References",
        "items": [
            "Chao et al. Jailbreaking Black Box Large Language Models in Twenty Queries. arXiv:2310.08419.",
            "Mehrotra et al. Tree of Attacks: Jailbreaking Black-Box LLMs Automatically. arXiv:2312.02119.",
            "Zou et al. Universal and Transferable Adversarial Attacks on Aligned Language Models. arXiv:2307.15043.",
            "Inan et al. Llama Guard: LLM-based Input-Output Safeguard for Human-AI Conversations. arXiv:2312.06674.",
        ],
    },
]


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_textbox(slide, x, y, w, h, text, size=18, color=TEXT, bold=False,
                align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return box


def add_header(slide, title):
    add_textbox(slide, 0.6, 0.35, 12.1, 0.45, title, size=25, bold=True)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.92), Inches(12.1), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER
    line.line.fill.background()


def add_footer(slide, idx):
    add_textbox(
        slide, 0.6, 7.08, 9.4, 0.22,
        "Detection-Only Defenses Phase 1  |  AdvBench, PAIR, TAP, detector-only defenses",
        size=8.5, color=MUTED,
    )
    add_textbox(slide, 12.2, 7.08, 0.5, 0.22, str(idx), size=8.5,
                color=MUTED, align=PP_ALIGN.RIGHT)


def add_panel(slide, x, y, w, h, fill=PANEL, line=BORDER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_bullets(slide, bullets, x=0.8, y=1.25, w=11.8, h=5.7, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = TEXT
    return box


def add_title_slide(prs, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_textbox(slide, 0.8, 2.05, 11.8, 0.7, slide_data["title"],
                size=36, color=TEXT, bold=True)
    add_textbox(slide, 0.84, 2.92, 11.2, 0.45, slide_data["subtitle"],
                size=19, color=MUTED)
    add_panel(slide, 0.84, 3.75, 4.0, 0.55, fill=RGBColor(0x12, 0x2C, 0x24),
              line=GREEN)
    add_textbox(slide, 1.05, 3.9, 3.5, 0.2, "Adaptive attacks vs. detector-only defenses",
                size=11, color=GREEN, bold=True)
    add_textbox(slide, 0.84, 6.7, 11.4, 0.3, slide_data["footer"], size=11,
                color=MUTED)
    return slide


def add_standard_slide(prs, idx, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, slide_data["title"])
    add_bullets(slide, slide_data["bullets"])
    add_footer(slide, idx)
    return slide


def add_agenda_slide(prs, idx, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, slide_data["title"])
    for i, item in enumerate(slide_data["items"], start=1):
        y = 1.25 + (i - 1) * 0.78
        add_panel(slide, 0.85, y, 0.48, 0.48, fill=RGBColor(0x1E, 0x3A, 0x5F),
                  line=BLUE)
        add_textbox(slide, 1.0, y + 0.12, 0.18, 0.15, str(i), size=12,
                    color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, 1.55, y + 0.08, 10.4, 0.32, item, size=18,
                    color=TEXT)
    add_footer(slide, idx)
    return slide


def add_two_col_slide(prs, idx, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, slide_data["title"])
    add_panel(slide, 0.75, 1.25, 5.8, 5.25)
    add_panel(slide, 6.8, 1.25, 5.8, 5.25)
    add_textbox(slide, 1.05, 1.55, 5.2, 0.35, slide_data["left_title"],
                size=19, color=GREEN, bold=True)
    add_textbox(slide, 7.1, 1.55, 5.2, 0.35, slide_data["right_title"],
                size=19, color=BLUE, bold=True)
    add_bullets(slide, slide_data["left"], x=1.05, y=2.1, w=5.05, h=3.8, size=15)
    add_bullets(slide, slide_data["right"], x=7.1, y=2.1, w=5.05, h=3.8, size=15)
    add_footer(slide, idx)
    return slide


def add_diagram_slide(prs, idx, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, slide_data["title"])
    x0, y0 = 0.65, 2.05
    box_w, gap = 1.75, 0.35
    colors = [BLUE, GREEN, AMBER, ORANGE, ROSE, MUTED]
    for i, (label, sub) in enumerate(slide_data["steps"]):
        x = x0 + i * (box_w + gap)
        add_panel(slide, x, y0, box_w, 1.05, fill=RGBColor(0x12, 0x14, 0x1C),
                  line=colors[i])
        add_textbox(slide, x + 0.12, y0 + 0.22, box_w - 0.24, 0.22, label,
                    size=14, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.12, y0 + 0.58, box_w - 0.24, 0.18, sub,
                    size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < len(slide_data["steps"]) - 1:
            add_textbox(slide, x + box_w + 0.05, y0 + 0.36, 0.25, 0.2, "->",
                        size=18, color=MUTED, bold=True)
    add_panel(slide, 2.3, 4.05, 8.7, 0.95, fill=RGBColor(0x24, 0x1B, 0x10),
              line=AMBER)
    add_textbox(slide, 2.55, 4.35, 8.2, 0.22, slide_data["caption"], size=15,
                color=TEXT, align=PP_ALIGN.CENTER)
    add_footer(slide, idx)
    return slide


def add_metrics_slide(prs, idx, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, slide_data["title"])
    positions = [(0.8, 1.45), (6.9, 1.45), (0.8, 4.05), (6.9, 4.05)]
    colors = [GREEN, BLUE, AMBER, ROSE]
    for (name, formula), (x, y), color in zip(slide_data["metrics"], positions, colors):
        add_panel(slide, x, y, 5.55, 1.75, fill=RGBColor(0x12, 0x14, 0x1C), line=color)
        add_textbox(slide, x + 0.25, y + 0.28, 5.0, 0.3, name, size=22,
                    color=color, bold=True)
        add_textbox(slide, x + 0.25, y + 0.95, 5.0, 0.3, formula, size=15,
                    color=TEXT)
    add_footer(slide, idx)
    return slide


def add_references_slide(prs, idx, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, slide_data["title"])
    add_bullets(slide, slide_data["items"], size=14)
    add_footer(slide, idx)
    return slide


def build_deck():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Missing source briefing: {SOURCE}")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_no = 0
    for data in SLIDES:
        kind = data.get("kind", "standard")
        if kind == "title":
            add_title_slide(prs, data)
        elif kind == "agenda":
            slide_no += 1
            add_agenda_slide(prs, slide_no, data)
        elif kind == "two_col":
            slide_no += 1
            add_two_col_slide(prs, slide_no, data)
        elif kind == "diagram":
            slide_no += 1
            add_diagram_slide(prs, slide_no, data)
        elif kind == "metrics":
            slide_no += 1
            add_metrics_slide(prs, slide_no, data)
        elif kind == "references":
            slide_no += 1
            add_references_slide(prs, slide_no, data)
        else:
            slide_no += 1
            add_standard_slide(prs, slide_no, data)

    os.makedirs(OUTPUT.parent, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build_deck()
